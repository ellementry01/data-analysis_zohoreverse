"""File parsing, header detection, type inference and coercion."""
from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook

from app.services.ingest import IngestError, detect_header_row, frame_from_grid, preview_sheet, read_grids
from app.services.typeinfer import coerce, infer_type, parse_dates

from conftest import api


def _write(tmp_path, name: str, data: bytes | str):
    p = tmp_path / name
    p.write_bytes(data.encode("utf-8") if isinstance(data, str) else data)
    return str(p)


def _preview(tmp_path, name: str, data, header_row=None) -> dict:
    sheets = read_grids(_write(tmp_path, name, data), name)
    return preview_sheet(sheets[0], header_row)


def _col(prev: dict, source: str) -> dict:
    return next(c for c in prev["columns"] if c["source_name"] == source)


# ----------------------------------------------------------------------------------------------- CSV parsing

def test_utf8_bom_is_stripped(tmp_path):
    prev = _preview(tmp_path, "bom.csv", b"\xef\xbb\xbfid,name\n1,a\n2,b\n")
    assert [c["source_name"] for c in prev["columns"]] == ["id", "name"]
    assert prev["columns"][0]["name"] == "id"


def test_cp1252_fallback(tmp_path):
    raw = "city,amount\nCafé Town,10\nNaïve,20\n".encode("cp1252")
    prev = _preview(tmp_path, "latin.csv", raw)
    assert prev["rows"][0][0] == "Café Town"
    assert prev["rows"][1][0] == "Naïve"


def test_semicolon_delimiter_sniffed(tmp_path):
    prev = _preview(tmp_path, "semi.csv", "region;qty;price\nNorth;1;10\nSouth;2;20\nEast;3;30\n")
    assert [c["source_name"] for c in prev["columns"]] == ["region", "qty", "price"]
    assert _col(prev, "qty")["inferred_type"] == "integer"
    assert prev["row_count"] == 3


def test_tab_delimiter_tsv_and_txt(tmp_path):
    content = "region\tqty\nNorth\t1\nSouth\t2\n"
    for name in ("data.tsv", "data.txt"):
        prev = _preview(tmp_path, name, content)
        assert [c["source_name"] for c in prev["columns"]] == ["region", "qty"], name
        assert prev["rows"] == [["North", "1"], ["South", "2"]]


def test_empty_and_unsupported_files(tmp_path):
    with pytest.raises(IngestError):
        read_grids(_write(tmp_path, "empty.csv", b"  \n"), "empty.csv")
    with pytest.raises(IngestError):
        read_grids(_write(tmp_path, "x.pdf", b"%PDF"), "x.pdf")


def test_header_detection_skips_title_rows(tmp_path):
    content = ("Sales report - exported from ERP,,\n"
               "Generated 2026-10-16,,\n"
               ",,\n"
               "Date,Region,Amount\n"
               "2026-01-01,North,10\n"
               "2026-01-02,South,20\n")
    prev = _preview(tmp_path, "title.csv", content)
    assert prev["header_row"] == 3
    assert [c["source_name"] for c in prev["columns"]] == ["Date", "Region", "Amount"]
    assert prev["row_count"] == 2


def test_blank_and_duplicate_headers(tmp_path):
    content = "Amount,,Amount,Name\n1,x,3,a\n2,y,4,b\n"
    prev = _preview(tmp_path, "dups.csv", content)
    src = [c["source_name"] for c in prev["columns"]]
    assert src == ["Amount", "Column 2", "Amount (2)", "Name"]
    names = [c["name"] for c in prev["columns"]]
    assert len(set(names)) == 4
    assert any(w["code"] == "missing_header" for w in prev["warnings"])


def test_total_footer_and_repeated_header_and_blank_rows_dropped(tmp_path):
    content = ("Region,Qty\n"
               "North,1\n"
               ",\n"
               "South,2\n"
               "Region,Qty\n"
               "East,3\n"
               "Total,6\n"
               "GRAND TOTAL,6\n")
    prev = _preview(tmp_path, "footer.csv", content)
    assert [r[0] for r in prev["rows"]] == ["North", "South", "East"]
    assert _col(prev, "Qty")["inferred_type"] == "integer"


def test_fully_empty_column_dropped(tmp_path):
    # a column with no header and no values (trailing comma exports) disappears
    prev = _preview(tmp_path, "emptycol.csv", "a,,c\n1,,x\n2,,y\n")
    assert [c["source_name"] for c in prev["columns"]] == ["a", "c"]
    # a headed but empty column is kept (all values empty)
    prev = _preview(tmp_path, "emptycol2.csv", "a,b,c\n1,,x\n2,,y\n")
    assert _col(prev, "b")["null_count"] == 2


# ----------------------------------------------------------------------------------------------- type inference

def _s(values):
    return pd.Series(values, dtype=object)


def test_infer_integers():
    assert infer_type(_s(["1", "2", "30", "-4"]))[0] == "integer"


def test_infer_decimals_with_currency_and_negatives():
    s = _s(["₹1,234.50", "Rs. 2,000", "(500)", "12%", "3.25"])
    t, fails = infer_type(s, "amount")
    assert t == "decimal" and fails == 0
    conv, failed, _ = coerce(s, "decimal")
    assert failed == 0
    assert conv.tolist() == pytest.approx([1234.5, 2000.0, -500.0, 0.12, 3.25])


def test_percent_column_is_decimal_fraction():
    s = _s(["12%", "15%", "100%"])
    assert infer_type(s)[0] == "decimal"
    assert coerce(s, "decimal")[0].tolist() == pytest.approx([0.12, 0.15, 1.0])


def test_rupee_integers_are_integers():
    s = _s(["₹1,000", "Rs.2,500", "INR 300"])
    assert infer_type(s)[0] == "integer"
    assert coerce(s, "integer")[0].tolist() == [1000, 2500, 300]


def test_leading_zero_ids_stay_text():
    s = _s(["00123", "00456", "01234", "00999"])
    assert infer_type(s, "zip")[0] == "text"
    assert coerce(s, "text")[0].tolist() == ["00123", "00456", "01234", "00999"]


def test_booleans():
    assert infer_type(_s(["Yes", "No", "yes", "NO"]))[0] == "boolean"
    assert infer_type(_s(["true", "false", "TRUE"]))[0] == "boolean"
    assert coerce(_s(["Yes", "no", "maybe"]), "boolean")[0].tolist() == [True, False, None]
    # a 0/1 column is a number, not a boolean
    assert infer_type(_s(["1", "0", "1"]))[0] == "integer"


def test_dates_iso_and_dayfirst_mixed_in_one_column():
    s = _s(["2024-01-15", "13/02/2024", "05/03/2024", "2024-04-01", "2024-06-30", "01/07/2024"])
    t, fails = infer_type(s, "invoice_date")
    assert t == "date" and fails == 0
    parsed = parse_dates(s)
    assert [d.date() for d in parsed] == [date(2024, 1, 15), date(2024, 2, 13), date(2024, 3, 5),
                                          date(2024, 4, 1), date(2024, 6, 30), date(2024, 7, 1)]


def test_dayfirst_only_ambiguous_dates():
    parsed = parse_dates(_s(["05/03/2024", "06/03/2024", "25/12/2024"]))
    assert [d.date() for d in parsed] == [date(2024, 3, 5), date(2024, 3, 6), date(2024, 12, 25)]


def test_datetimes():
    s = _s(["2024-01-15 10:30:00", "2024-01-16 08:00:00", "2024-01-17 23:59:59"])
    assert infer_type(s, "created_at")[0] == "datetime"
    conv, failed, _ = coerce(s, "datetime")
    assert failed == 0 and conv.iloc[0] == pd.Timestamp("2024-01-15 10:30:00")


def test_coercion_failure_counts():
    s = _s([str(i) for i in range(1, 11)] + ["abc"])
    t, fails = infer_type(s, "qty")
    assert t == "integer" and fails == 1
    conv, failed, bad = coerce(s, "integer")
    assert failed == 1 and bad == ["abc"]
    assert conv.isna().sum() == 1


def test_null_markers_are_empty(tmp_path):
    prev = _preview(tmp_path, "nulls.csv", "id,v\n1,N/A\n2,-\n3,5\n4,null\n5,7\n")
    c = _col(prev, "v")
    assert c["null_count"] == 3
    assert c["inferred_type"] == "integer"


def test_text_column_variants_reported(tmp_path):
    prev = _preview(tmp_path, "variants.csv", "id,channel\n1,Amazon\n2,amazon\n3, Amazon\n4,Shopify\n")
    c = _col(prev, "channel")
    assert c["inconsistent_variants"]
    assert any(w["code"] == "inconsistent_values" for w in prev["warnings"])


# ----------------------------------------------------------------------------------------------- Excel

def test_excel_with_title_rows_typed_cells_and_two_sheets(tmp_path):
    wb = Workbook()
    sh = wb.active
    sh.title = "Product Master"
    sh.append(["Product Master - exported from ERP"])
    sh.append(["Generated 2026-10-16"])
    sh.append([])
    sh.append(["SKU", "Product Name", "Price", "Launch Date", "Active"])
    sh.append(["0001", "Kadai", 1299.5, datetime(2024, 1, 15), "Yes"])
    sh.append(["0002", "Tawa", 799, datetime(2024, 2, 1), "No"])
    sh.append(["0003", "Jug", "Rs. 1,050", datetime(2024, 3, 10), "Yes"])
    sh.append(["Total", None, 3148.5, None, None])
    notes = wb.create_sheet("Notes")
    notes.append(["Notes"])
    notes.append(["MRP includes GST"])
    p = tmp_path / "products.xlsx"
    wb.save(p)

    sheets = read_grids(str(p), "products.xlsx")
    assert [s.name for s in sheets] == ["Product Master", "Notes"]
    prev = preview_sheet(sheets[0])
    assert [c["source_name"] for c in prev["columns"]] == ["SKU", "Product Name", "Price", "Launch Date", "Active"]
    types = {c["source_name"]: c["inferred_type"] for c in prev["columns"]}
    assert types == {"SKU": "text", "Product Name": "text", "Price": "decimal", "Launch Date": "date", "Active": "boolean"}
    assert prev["row_count"] == 3  # the Total footer is dropped


def test_excel_upload_and_import_via_api(client, workspace, tmp_path):
    wb = Workbook()
    sh = wb.active
    sh.title = "Data"
    sh.append(["Report title"])
    sh.append(["Date", "Qty", "Amount"])
    sh.append([datetime(2026, 1, 1), 1, "₹100.50"])
    sh.append([datetime(2026, 1, 2), 2, "(20)"])
    sh.append([datetime(2026, 1, 3), "x", "30"])
    p = tmp_path / "book.xlsx"
    wb.save(p)
    ws = workspace["id"]
    r = client.post(api(ws, "/uploads"), files={"file": ("book.xlsx", p.read_bytes())})
    assert r.status_code == 200, r.text
    prev = r.json()
    assert prev["header_row"] == 1
    qty = next(c for c in prev["columns"] if c["source_name"] == "Qty")
    r = client.post(api(ws, "/imports"), json={"upload_id": prev["upload_id"], "sheet": prev["sheet"], "header_row": 1,
                                               "columns": [{"source_name": c["source_name"], "name": c["name"],
                                                            "data_type": "integer" if c["source_name"] == "Qty" else c["inferred_type"]}
                                                           for c in prev["columns"]]})
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["report"]["conversion_failures"]["qty"]["count"] == 1
    assert body["report"]["conversion_failures"]["qty"]["samples"] == ["x"]
    assert qty["parse_failures"] in (0, 1)
    ds = body["dataset"]
    rows = client.get(api(ws, f"/datasets/{ds['id']}/rows"), params={"sort": "date_"}).json()
    by = dict(zip(rows["columns"], zip(*rows["rows"])))
    assert list(by["amount"]) == pytest.approx([100.5, -20.0, 30.0])
    assert list(by["qty"]) == [1, 2, None]
    assert by["date_"][0].startswith("2026-01-01")


def test_header_row_override_and_unknown_sheet(client, workspace):
    ws = workspace["id"]
    r = client.post(api(ws, "/uploads"), files={"file": ("h.csv", b"junk,junk2\na,b\n1,2\n")})
    prev = r.json()
    r = client.get(api(ws, f"/uploads/{prev['upload_id']}/preview"), params={"header_row": 1})
    assert r.status_code == 200
    assert [c["source_name"] for c in r.json()["columns"]] == ["a", "b"]
    r = client.get(api(ws, f"/uploads/{prev['upload_id']}/preview"), params={"sheet": "nope"})
    assert r.status_code == 400
    r = client.get(api(ws, "/uploads/zzzz/preview"))
    assert r.status_code == 400


def test_detect_header_row_unit():
    grid = pd.DataFrame([["Title", None, None], [None, None, None], ["a", "b", "c"], ["1", "2", "3"]], dtype=object)
    assert detect_header_row(grid) == 2
    df, headers = frame_from_grid(grid, 2)
    assert headers == ["a", "b", "c"] and len(df) == 1
