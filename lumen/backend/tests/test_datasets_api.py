"""Datasets through the HTTP API: import, browse, prepare, model, export, roles and database connectors."""
import io
import json

import pytest
from sqlalchemy import text

from conftest import api, upload_csv

SALES_V1 = """order_id,order_date,region,product,qty,price,notes
1,2024-01-05,North,Widget,2,10.50,a
2,2024-01-20,north,Gadget,1,20.00,
3,2024-02-03,South,Widget,5,10.50,b
3,2024-02-03,South,Widget,5,10.50,b
4,2024-02-15,East,Gizmo,,7.25,c
5,2024-03-01,NORTH,Gizmo,3,7.25,d
"""

SALES_V2 = """order_id,order_date,region,product,qty,price,notes
10,2024-04-01,north,Gizmo,,7.25,x
11,2024-04-02,East,Widget,1,10.50,y
11,2024-04-02,East,Widget,1,10.50,y
12,2024-04-03,south,Gadget,2,20.00,z
13,2024-04-04,NORTH,Widget,4,10.50,w
"""

SALES_V3_EXTRA_COL = """order_id,order_date,region,product,qty,price,discount
20,2024-05-01,West,Widget,1,10.50,5
21,2024-05-02,West,Gadget,2,20.00,0
"""

COLUMN_SPEC = [
    {"source_name": "order_id", "name": "order_id", "data_type": "text"},
    {"source_name": "order_date", "name": "order_date", "data_type": "date"},
    {"source_name": "region", "name": "region", "data_type": "text"},
    {"source_name": "product", "name": "product", "data_type": "text"},
    {"source_name": "qty", "name": "qty", "data_type": "integer"},
    {"source_name": "price", "name": "price", "display_name": "Unit price", "data_type": "decimal"},
    {"source_name": "notes", "name": "notes", "data_type": "text", "include": False},
]


def _upload(client, ws, filename, content) -> dict:
    r = client.post(api(ws, "/uploads"), files={"file": (filename, content.encode())})
    assert r.status_code == 200, r.text
    return r.json()


def _import_sales(client, ws, content=SALES_V1, **kw) -> dict:
    prev = _upload(client, ws, "sales.csv", content)
    body = {"upload_id": prev["upload_id"], "sheet": prev["sheet"], "header_row": prev["header_row"], "name": "Sales",
            "columns": COLUMN_SPEC, **kw}
    r = client.post(api(ws, "/imports"), json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _rows(client, ws, ds_id, **params) -> dict:
    r = client.get(api(ws, f"/datasets/{ds_id}/rows"), params={"page_size": 500, **params})
    assert r.status_code == 200, r.text
    return r.json()


def _records(client, ws, ds_id, sort="order_id") -> list[dict]:
    d = _rows(client, ws, ds_id, sort=sort)
    return [dict(zip(d["columns"], r)) for r in d["rows"]]


def _dataset(client, ws, ds_id) -> dict:
    r = client.get(api(ws, f"/datasets/{ds_id}"))
    assert r.status_code == 200, r.text
    return r.json()


# ----------------------------------------------------------------------------------------------- import

def test_preview_reports_types_and_warnings(client, workspace):
    prev = _upload(client, workspace["id"], "sales.csv", SALES_V1)
    types = {c["source_name"]: c["inferred_type"] for c in prev["columns"]}
    assert types == {"order_id": "integer", "order_date": "date", "region": "text", "product": "text", "qty": "integer",
                     "price": "decimal", "notes": "text"}
    codes = {w["code"] for w in prev["warnings"]}
    assert {"duplicates", "inconsistent_values"} <= codes
    assert prev["row_count"] == 6 and len(prev["rows"]) == 6


def test_import_with_type_overrides_and_excluded_columns(client, workspace):
    ws = workspace["id"]
    res = _import_sales(client, ws)
    ds = res["dataset"]
    assert res["report"]["rows_loaded"] == 6
    assert ds["kind"] == "file" and ds["slug"] == "sales" and ds["row_count"] == 6
    cols = {c["name"]: c for c in ds["columns"]}
    assert "notes" not in cols
    assert cols["order_id"]["data_type"] == "text"
    assert cols["price"]["display_name"] == "Unit price"
    assert ds["default_date_column"] == "order_date"
    rec = _records(client, ws, ds["id"])
    assert rec[0]["order_id"] == "1" and rec[0]["order_date"] == "2024-01-05" and rec[0]["price"] == 10.5
    # listed in the workspace
    assert [d["slug"] for d in client.get(api(ws, "/datasets")).json()] == ["sales"]


def test_import_validation_errors(client, workspace):
    ws = workspace["id"]
    prev = _upload(client, ws, "sales.csv", SALES_V1)
    base = {"upload_id": prev["upload_id"], "sheet": prev["sheet"], "header_row": 0}
    r = client.post(api(ws, "/imports"), json={**base, "columns": [{**c, "include": False} for c in COLUMN_SPEC]})
    assert r.status_code == 400
    r = client.post(api(ws, "/imports"), json={**base, "columns": [{**c, "name": "same"} for c in COLUMN_SPEC]})
    assert r.status_code == 400
    r = client.post(api(ws, "/imports"), json={**base, "columns": [{**COLUMN_SPEC[0], "data_type": "money"}]})
    assert r.status_code == 400
    r = client.post(api(ws, "/imports"), json={**base, "mode": "append", "dataset_id": 999999})
    assert r.status_code == 400


def test_rows_paging_search_sort_and_filters(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    page = client.get(api(ws, f"/datasets/{ds['id']}/rows"), params={"page": 2, "page_size": 2, "sort": "order_date"}).json()
    assert page["total"] == 6 and len(page["rows"]) == 2
    assert [r[0] for r in page["rows"]] == ["4", "5"]
    found = _rows(client, ws, ds["id"], search="gadget")
    assert found["total"] == 1 and found["rows"][0][3] == "Gadget"
    by_qty = _rows(client, ws, ds["id"], sort="qty", desc=True)
    qtys = [r[4] for r in by_qty["rows"]]
    assert qtys[0] == 5 and qtys[-1] is None  # NULLS LAST
    flt = _rows(client, ws, ds["id"], filters=json.dumps([{"column": "region", "op": "in", "value": ["north"]}]))
    assert flt["total"] == 3  # case-insensitive membership
    flt = _rows(client, ws, ds["id"], filters=json.dumps([{"column": "price", "op": "between", "value": [8, 12]}]))
    assert flt["total"] == 3
    # unknown sort column is ignored, bad filter JSON is a 400
    assert _rows(client, ws, ds["id"], sort="nope")["total"] == 6
    assert client.get(api(ws, f"/datasets/{ds['id']}/rows"), params={"filters": "{bad"}).status_code == 400
    assert client.get(api(ws, "/datasets/999999/rows")).status_code == 404


# ----------------------------------------------------------------------------------------------- columns

def test_rename_hide_and_change_type(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.patch(api(ws, f"/datasets/{ds['id']}/columns/qty"), json={"display_name": "Quantity", "description": "units"})
    assert r.status_code == 200, r.text
    col = next(c for c in r.json()["columns"] if c["name"] == "qty")
    assert col["display_name"] == "Quantity" and col["description"] == "units"

    r = client.patch(api(ws, f"/datasets/{ds['id']}/columns/order_id"), json={"data_type": "integer"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert next(c for c in out["columns"] if c["name"] == "order_id")["data_type"] == "integer"
    assert out["transforms"][-1]["kind"] == "cast_type"
    assert [x["order_id"] for x in _records(client, ws, ds["id"])] == [1, 2, 3, 3, 4, 5]

    # a cast that cannot parse leaves values empty and reports it
    r = client.patch(api(ws, f"/datasets/{ds['id']}/columns/region"), json={"data_type": "date"})
    assert r.status_code == 200
    assert "6 values could not be converted" in r.json()["transforms"][-1]["last_result"]

    r = client.patch(api(ws, f"/datasets/{ds['id']}/columns/product"), json={"hidden": True})
    assert next(c for c in r.json()["columns"] if c["name"] == "product")["hidden"] is True
    model = client.get(api(ws, "/model")).json()
    assert "sales.product" not in {f["ref"] for d in model["datasets"] for f in d["fields"]}

    assert client.patch(api(ws, f"/datasets/{ds['id']}/columns/qty"), json={"data_type": "money"}).status_code == 400
    assert client.patch(api(ws, f"/datasets/{ds['id']}/columns/nope"), json={"display_name": "x"}).status_code == 400


def test_dataset_patch(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.patch(api(ws, f"/datasets/{ds['id']}"), json={"name": "Sales 2024", "description": "d"})
    assert r.json()["name"] == "Sales 2024" and r.json()["slug"] == "sales"
    assert client.patch(api(ws, f"/datasets/{ds['id']}"), json={"default_date_column": "region"}).status_code == 400


def test_derived_columns(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.post(api(ws, f"/datasets/{ds['id']}/columns"), json={"display_name": "Revenue", "expression": "qty * price"})
    assert r.status_code == 201, r.text
    col = next(c for c in r.json()["columns"] if c["name"] == "revenue")
    assert col["is_derived"] and col["data_type"] == "decimal"
    rec = _records(client, ws, ds["id"])
    assert rec[0]["revenue"] == 21.0 and rec[4]["revenue"] is None

    r = client.post(api(ws, f"/datasets/{ds['id']}/columns"),
                    json={"display_name": "Size", "expression": "CASE WHEN qty >= 3 THEN 'big' ELSE 'small' END", "data_type": "text"})
    assert r.status_code == 201, r.text
    assert _records(client, ws, ds["id"])[2]["size"] == "big"

    for expr in ("qty * nope", "SUM(qty)", "(SELECT 1)", "pg_sleep(1)", "region * 2"):
        r = client.post(api(ws, f"/datasets/{ds['id']}/columns"), json={"display_name": "Bad", "expression": expr})
        assert r.status_code == 400, expr
    assert client.post(api(ws, f"/datasets/{ds['id']}/columns"), json={"display_name": "Revenue", "expression": "qty"}).status_code == 400
    # the dataset is unharmed by the failed formulas
    d = _dataset(client, ws, ds["id"])
    assert "bad" not in {c["name"] for c in d["columns"]}
    assert _rows(client, ws, ds["id"])["total"] == 6

    # edit and delete the derived column
    r = client.patch(api(ws, f"/datasets/{ds['id']}/columns/revenue"), json={"expression": "qty * price * 2"})
    assert r.status_code == 200, r.text
    assert _records(client, ws, ds["id"])[0]["revenue"] == 42.0
    r = client.delete(api(ws, f"/datasets/{ds['id']}/columns/revenue"))
    assert r.status_code == 200
    assert "revenue" not in _rows(client, ws, ds["id"])["columns"]
    # deleting a physical column records a drop_column step
    r = client.delete(api(ws, f"/datasets/{ds['id']}/columns/product"))
    assert r.status_code == 200 and r.json()["transforms"][-1]["kind"] == "drop_column"
    assert "product" not in _rows(client, ws, ds["id"])["columns"]


# ----------------------------------------------------------------------------------------------- transforms

def _step(client, ws, ds_id, kind, params) -> dict:
    r = client.post(api(ws, f"/datasets/{ds_id}/transforms"), json={"kind": kind, "params": params})
    assert r.status_code == 200, r.text
    return r.json()


def test_transforms_and_replay_after_replace(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    ds_id = ds["id"]
    client.post(api(ws, f"/datasets/{ds_id}/columns"), json={"display_name": "Revenue", "expression": "qty * price"})
    assert "values trimmed" in _step(client, ws, ds_id, "trim_whitespace", {})["message"]
    assert _step(client, ws, ds_id, "standardize_values", {"column": "region", "mapping": {"North": ["north", "NORTH"]}})["message"] == "2 values standardized"
    assert _step(client, ws, ds_id, "remove_duplicates", {})["message"] == "1 duplicate rows removed"
    assert _step(client, ws, ds_id, "fill_nulls", {"column": "qty", "value": 0})["message"] == "1 empty values filled"
    assert _step(client, ws, ds_id, "delete_rows", {"filters": [{"column": "region", "op": "eq", "value": "East"}]})["message"] == "1 rows deleted"
    out = _step(client, ws, ds_id, "replace_values", {"column": "product", "find": "Gizmo", "replace": "Gizmo Pro"})
    assert out["message"] == "1 values replaced"
    assert [t["kind"] for t in out["dataset"]["transforms"]] == [
        "trim_whitespace", "standardize_values", "remove_duplicates", "fill_nulls", "delete_rows", "replace_values"]
    rec = _records(client, ws, ds_id)
    assert [(r["order_id"], r["region"], r["product"], r["qty"]) for r in rec] == [
        ("1", "North", "Widget", 2), ("2", "North", "Gadget", 1), ("3", "South", "Widget", 5), ("5", "North", "Gizmo Pro", 3)]
    assert _dataset(client, ws, ds_id)["row_count"] == 4

    # bad steps are rejected and not recorded
    r = client.post(api(ws, f"/datasets/{ds_id}/transforms"), json={"kind": "explode", "params": {}})
    assert r.status_code == 400
    r = client.post(api(ws, f"/datasets/{ds_id}/transforms"), json={"kind": "fill_nulls", "params": {"column": "nope", "value": 1}})
    assert r.status_code == 400
    r = client.post(api(ws, f"/datasets/{ds_id}/transforms"), json={"kind": "delete_rows", "params": {"filters": []}})
    assert r.status_code == 400
    assert len(_dataset(client, ws, ds_id)["transforms"]) == 6

    # replacing the data replays every recorded step on the new rows
    res = _import_sales(client, ws, SALES_V2, mode="replace", dataset_id=ds_id)
    assert res["report"]["transform_messages"] == []
    rec = _records(client, ws, ds_id)
    assert [(r["order_id"], r["region"], r["product"], r["qty"]) for r in rec] == [
        ("10", "North", "Gizmo Pro", 0), ("12", "south", "Gadget", 2), ("13", "North", "Widget", 4)]
    assert rec[2]["revenue"] == 42.0  # derived column survives the replace
    d = _dataset(client, ws, ds_id)
    assert d["row_count"] == 3 and len(d["transforms"]) == 6

    # removing a step stops it from being replayed
    step = next(t for t in d["transforms"] if t["kind"] == "delete_rows")
    r = client.delete(api(ws, f"/datasets/{ds_id}/transforms/{step['id']}"))
    assert r.status_code == 200 and len(r.json()["transforms"]) == 5
    _import_sales(client, ws, SALES_V2, mode="replace", dataset_id=ds_id)
    assert [r["order_id"] for r in _records(client, ws, ds_id)] == ["10", "11", "12", "13"]


def test_cast_type_is_replayed_after_replace(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    client.patch(api(ws, f"/datasets/{ds['id']}/columns/order_id"), json={"data_type": "integer"})
    _import_sales(client, ws, SALES_V2, mode="replace", dataset_id=ds["id"])
    d = _dataset(client, ws, ds["id"])
    assert next(c for c in d["columns"] if c["name"] == "order_id")["data_type"] == "integer"
    assert [r["order_id"] for r in _records(client, ws, ds["id"])] == [10, 11, 11, 12, 13]


def test_change_case_and_contains_replace(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    _step(client, ws, ds["id"], "change_case", {"column": "region", "mode": "upper"})
    _step(client, ws, ds["id"], "replace_values", {"column": "product", "find": "dg", "replace": "DG", "match": "contains"})
    rec = _records(client, ws, ds["id"])
    assert {r["region"] for r in rec} == {"NORTH", "SOUTH", "EAST"}
    assert rec[1]["product"] == "GaDGet"


def test_append_adds_rows_and_new_columns(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    prev = _upload(client, ws, "more.csv", SALES_V3_EXTRA_COL)
    r = client.post(api(ws, "/imports"), json={"upload_id": prev["upload_id"], "sheet": prev["sheet"], "header_row": 0,
                                               "mode": "append", "dataset_id": ds["id"]})
    assert r.status_code == 201, r.text
    out = r.json()
    assert out["report"]["rows_loaded"] == 2
    assert out["dataset"]["row_count"] == 8
    cols = {c["name"]: c for c in out["dataset"]["columns"]}
    assert cols["discount"]["data_type"] == "integer"
    rec = _rows(client, ws, ds["id"], sort="order_date")
    recs = [dict(zip(rec["columns"], x)) for x in rec["rows"]]
    assert recs[0]["discount"] is None and recs[-1]["discount"] == 0 and recs[-1]["region"] == "West"
    # order_id stays text (type of the existing dataset wins)
    assert recs[-1]["order_id"] == "21"


def test_query_table_cannot_receive_rows(client, workspace):
    ws = workspace["id"]
    _import_sales(client, ws)
    qt = client.post(api(ws, "/query-tables"), json={"name": "Q", "sql": "SELECT region FROM sales"}).json()
    prev = _upload(client, ws, "more.csv", SALES_V1)
    r = client.post(api(ws, "/imports"), json={"upload_id": prev["upload_id"], "header_row": 0, "mode": "append", "dataset_id": qt["id"]})
    assert r.status_code == 400


def test_persistence_across_new_client_and_engine(client, workspace):
    from fastapi.testclient import TestClient

    from app.db import reset_engines
    from app.main import app

    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    _step(client, ws, ds["id"], "remove_duplicates", {})
    reset_engines()  # drop every pooled connection: the next client starts from what is on disk
    with TestClient(app) as c2:
        d = c2.get(api(ws, f"/datasets/{ds['id']}")).json()
        assert d["row_count"] == 5 and d["transforms"][0]["kind"] == "remove_duplicates"
        assert c2.get(api(ws, f"/datasets/{ds['id']}/rows")).json()["total"] == 5
        assert any(w["id"] == ws for w in c2.get("/api/workspaces").json())


# ----------------------------------------------------------------------------------------------- query tables

def test_query_tables(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.post(api(ws, "/query-tables"), json={
        "name": "Region totals", "sql": "SELECT region, SUM(qty) AS total_qty, COUNT(*) AS n FROM sales GROUP BY region"})
    assert r.status_code == 201, r.text
    qt = r.json()
    assert qt["kind"] == "query_table" and qt["source_info"]["depends_on"] == ["sales"]
    assert {c["name"]: c["data_type"] for c in qt["columns"]} == {"region": "text", "total_qty": "decimal", "n": "integer"}
    rec = {r[0]: r[1] for r in _rows(client, ws, qt["id"])["rows"]}
    assert rec["South"] == 10 and rec["East"] is None

    # a query table on top of a query table
    r = client.post(api(ws, "/query-tables"), json={"name": "Big regions", "sql": "SELECT region FROM region_totals WHERE n > 1"})
    assert r.status_code == 201, r.text
    qt2 = r.json()

    # dependent views survive a replace of the base dataset
    _import_sales(client, ws, SALES_V2, mode="replace", dataset_id=ds["id"])
    rec = {r[0]: r[1] for r in _rows(client, ws, qt["id"])["rows"]}
    assert rec["East"] == 2 and rec["NORTH"] == 4
    assert {r[0] for r in _rows(client, ws, qt2["id"])["rows"]} == {"East"}

    # ... and a derived-column change on the base (which rebuilds the base view)
    client.post(api(ws, f"/datasets/{ds['id']}/columns"), json={"display_name": "Revenue", "expression": "qty * price"})
    client.delete(api(ws, f"/datasets/{ds['id']}/columns/revenue"))
    assert _rows(client, ws, qt2["id"])["total"] == 1

    # unsafe or invalid SQL is rejected
    for sql in ("DELETE FROM sales", "SELECT * FROM lumen_meta.workspaces", "SELECT * FROM _raw_sales",
                "SELECT pg_sleep(1)", "SELECT nope FROM sales", "SELECT region, region FROM sales", "SELECT 1 + 1 FROM sales"):
        r = client.post(api(ws, "/query-tables"), json={"name": "x", "sql": sql})
        assert r.status_code == 400, sql

    # update the SQL
    r = client.put(api(ws, f"/query-tables/{qt['id']}"), json={"name": "Region totals",
                                                               "sql": "SELECT region, SUM(qty) AS total_qty, COUNT(*) AS n, MAX(price) AS max_price FROM sales GROUP BY region"})
    assert r.status_code == 200, r.text
    assert "max_price" in {c["name"] for c in r.json()["columns"]}
    assert _rows(client, ws, qt2["id"])["total"] == 1
    assert client.put(api(ws, f"/query-tables/{ds['id']}"), json={"name": "x", "sql": "SELECT 1 AS a"}).status_code == 400

    # transforms / derived columns do not apply to query tables
    assert client.post(api(ws, f"/datasets/{qt['id']}/transforms"), json={"kind": "trim_whitespace", "params": {}}).status_code == 400
    assert client.post(api(ws, f"/datasets/{qt['id']}/columns"), json={"display_name": "x", "expression": "n * 2"}).status_code == 400

    # deleting a dataset that a query table depends on is blocked
    r = client.delete(api(ws, f"/datasets/{ds['id']}"))
    assert r.status_code == 400 and "Region totals" in r.json()["detail"]
    r = client.delete(api(ws, f"/datasets/{qt['id']}"))
    assert r.status_code == 400  # Big regions depends on it
    assert client.delete(api(ws, f"/datasets/{qt2['id']}")).status_code == 204
    assert client.delete(api(ws, f"/datasets/{qt['id']}")).status_code == 204
    assert client.delete(api(ws, f"/datasets/{ds['id']}")).status_code == 204
    assert client.get(api(ws, "/datasets")).json() == []

    lineage = client.get(api(ws, "/lineage")).json()
    assert lineage["nodes"] == []


# ----------------------------------------------------------------------------------------------- relationships

ORDERS = """order_id,order_date,product_code,qty
1,2024-01-01,PC-1,2
2,2024-01-02,PC-2,1
3,2024-01-03,PC-1,4
4,2024-01-04,PC-9,1
"""
PRODUCTS = """product_code,product_name
PC-1,Widget
PC-2,Gadget
PC-3,Gizmo
"""
DUP_PRODUCTS = """product_code,label
PC-1,Widget
PC-1,Widget again
"""


def test_relationships(client, workspace):
    ws = workspace["id"]
    o = upload_csv(client, ws, "orders.csv", ORDERS, name="Orders")["dataset"]
    p = upload_csv(client, ws, "products.csv", PRODUCTS, name="Products")["dataset"]
    dp = upload_csv(client, ws, "dup.csv", DUP_PRODUCTS, name="Dup products")["dataset"]

    sugg = client.get(api(ws, "/relationships/suggestions")).json()
    pair = [(s["from_dataset_id"], s["from_column"], s["to_dataset_id"], s["to_column"]) for s in sugg]
    assert (o["id"], "product_code", p["id"], "product_code") in pair
    assert all(s["to_dataset_id"] != dp["id"] for s in sugg)  # non-unique key never suggested
    assert all(s["to_dataset_id"] != o["id"] for s in sugg)

    body = {"from_dataset_id": o["id"], "from_column": "product_code", "to_dataset_id": p["id"], "to_column": "product_code"}
    r = client.post(api(ws, "/relationships"), json=body)
    assert r.status_code == 201, r.text
    rel = r.json()
    assert rel["stats"] == {"from_values": 4, "matched": 3, "match_rate": 0.75, "lookup_unique": True}
    assert rel["cardinality"] == "many_to_one"
    assert client.post(api(ws, "/relationships"), json=body).status_code == 400  # one path per pair
    assert all(s["from_dataset_id"] != o["id"] for s in client.get(api(ws, "/relationships/suggestions")).json())

    r = client.post(api(ws, "/relationships"), json={**body, "to_dataset_id": dp["id"]})
    assert r.status_code == 400 and "duplicate" in r.json()["detail"]
    r = client.post(api(ws, "/relationships"), json={**body, "from_column": "order_date", "to_dataset_id": dp["id"]})
    assert r.status_code == 400 and "types differ" in r.json()["detail"]
    r = client.post(api(ws, "/relationships"), json={**body, "to_dataset_id": o["id"], "to_column": "product_code"})
    assert r.status_code == 400
    r = client.post(api(ws, "/relationships"), json={**body, "from_column": "nope"})
    assert r.status_code == 400

    rels = client.get(api(ws, "/relationships")).json()
    assert len(rels) == 1
    model = client.get(api(ws, "/model")).json()
    orders = next(d for d in model["datasets"] if d["slug"] == "orders")
    assert orders["lookups"] == ["products"]
    assert {"from": "orders.product_code", "to": "products.product_code"} in model["relationships"]

    # deleting a dataset removes its relationships
    assert client.delete(api(ws, f"/datasets/{p['id']}")).status_code == 204
    assert client.get(api(ws, "/relationships")).json() == []
    assert client.delete(api(ws, f"/relationships/{rels[0]['id']}")).status_code == 404


# ----------------------------------------------------------------------------------------------- export / profile / quality

def test_export_csv_and_xlsx(client, workspace):
    from openpyxl import load_workbook

    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.get(api(ws, f"/datasets/{ds['id']}/export"), params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="sales.csv"' in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "Order id,Order date,Region,Product,Qty,Unit price"
    assert len(lines) == 7 and lines[1].startswith("1,2024-01-05,North,Widget,2,10.5")

    r = client.get(api(ws, f"/datasets/{ds['id']}/export"), params={"format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][5] == "Unit price" and len(rows) == 7 and rows[1][4] == 2

    assert client.get(api(ws, f"/datasets/{ds['id']}/export"), params={"format": "pdf"}).status_code == 422


def test_profile(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    prof = client.get(api(ws, f"/datasets/{ds['id']}/profile")).json()
    assert prof["row_count"] == 6
    by = {c["name"]: c for c in prof["columns"]}
    assert by["qty"]["min"] == 1 and by["qty"]["max"] == 5 and by["qty"]["sum"] == 16 and by["qty"]["nulls"] == 1
    assert by["qty"]["median"] == 3
    assert sum(b["count"] for b in by["qty"]["histogram"]) == 5
    assert by["order_date"]["min"] == "2024-01-05" and by["order_date"]["max"] == "2024-03-01"
    assert by["region"]["distinct"] == 5
    assert by["product"]["top_values"][0] == {"value": "Widget", "count": 3}


def test_quality_warnings(client, workspace):
    ws = workspace["id"]
    content = SALES_V1 + "6,2024-03-05,South,Widget,-2,10.50,e\n"
    ds = _import_sales(client, ws, content)["dataset"]
    q = ds["quality"]
    codes = {(w["code"], w.get("column")) for w in q["warnings"]}
    assert ("duplicates", None) in codes
    assert ("missing", "qty") in codes
    assert ("inconsistent_values", "region") in codes
    assert ("negative", "qty") in codes
    assert ds["warning_count"] == len(q["warnings"])
    _step(client, ws, ds["id"], "remove_duplicates", {})
    q2 = client.post(api(ws, f"/datasets/{ds['id']}/quality")).json()
    assert "duplicates" not in {w["code"] for w in q2["warnings"]}


# ----------------------------------------------------------------------------------------------- roles

def test_roles(client, workspace):
    ws = workspace["id"]
    ds = _import_sales(client, ws)["dataset"]
    r = client.put(api(ws, "/members"), json={"user_email": "Viewer@Example.com", "role": "viewer"})
    assert r.status_code == 200
    client.put(api(ws, "/members"), json={"user_email": "editor@example.com", "role": "editor"})
    viewer = {"X-User-Email": "viewer@example.com"}
    editor = {"X-User-Email": "editor@example.com"}
    stranger = {"X-User-Email": "stranger@example.com"}

    assert client.get(api(ws, "/datasets"), headers=viewer).status_code == 200
    assert client.get(api(ws, f"/datasets/{ds['id']}/rows"), headers=viewer).status_code == 200
    r = client.post(api(ws, "/uploads"), files={"file": ("s.csv", SALES_V1.encode())}, headers=viewer)
    assert r.status_code == 403 and "editor" in r.json()["detail"]
    prev = _upload(client, ws, "sales.csv", SALES_V1)
    r = client.post(api(ws, "/imports"), json={"upload_id": prev["upload_id"], "header_row": 0}, headers=viewer)
    assert r.status_code == 403
    assert client.delete(api(ws, f"/datasets/{ds['id']}"), headers=viewer).status_code == 403
    assert client.post(api(ws, f"/datasets/{ds['id']}/transforms"), json={"kind": "trim_whitespace"}, headers=viewer).status_code == 403
    assert client.put(api(ws, "/members"), json={"user_email": "x@y.z", "role": "admin"}, headers=viewer).status_code == 403

    r = client.post(api(ws, "/imports"), json={"upload_id": prev["upload_id"], "header_row": 0, "name": "By editor"}, headers=editor)
    assert r.status_code == 201
    assert client.put(api(ws, "/members"), json={"user_email": "x@y.z", "role": "admin"}, headers=editor).status_code == 403
    assert client.delete(f"/api/workspaces/{ws}", headers=editor).status_code == 403

    for path in ("", "/datasets", f"/datasets/{ds['id']}/rows", "/model"):
        assert client.get(api(ws, path), headers=stranger).status_code == 403, path
    assert ws not in [w["id"] for w in client.get("/api/workspaces", headers=stranger).json()]
    assert [w["role"] for w in client.get("/api/workspaces", headers=viewer).json() if w["id"] == ws] == ["viewer"]

    # the last admin cannot be demoted or removed
    owner = client.get("/api/me").json()["email"]
    assert client.put(api(ws, "/members"), json={"user_email": owner, "role": "viewer"}).status_code == 400
    assert client.delete(api(ws, f"/members/{owner}")).status_code == 400
    assert client.delete(api(ws, "/members/viewer@example.com")).status_code == 200
    assert client.get(api(ws, "/datasets"), headers=viewer).status_code == 403
    assert client.get("/api/workspaces/999999/datasets").status_code == 404


# ----------------------------------------------------------------------------------------------- connectors

SRC_SCHEMA = "conn_src_test"


@pytest.fixture()
def source_table(client):
    from app.db import get_engine

    eng = get_engine()
    with eng.begin() as c:
        c.execute(text(f"DROP SCHEMA IF EXISTS {SRC_SCHEMA} CASCADE"))
        c.execute(text(f"CREATE SCHEMA {SRC_SCHEMA}"))
        c.execute(text(f"CREATE TABLE {SRC_SCHEMA}.items (id int PRIMARY KEY, name text, amount numeric(10,2), created date)"))
        c.execute(text(f"INSERT INTO {SRC_SCHEMA}.items VALUES (1, 'Kadai', 1299.50, '2026-01-01'), "
                       f"(2, 'Tawa', 799.00, '2026-02-01'), (3, 'Jug', 450.25, '2026-03-01')"))
        c.execute(text(f"CREATE VIEW {SRC_SCHEMA}.expensive AS SELECT * FROM {SRC_SCHEMA}.items WHERE amount > 500"))
    yield SRC_SCHEMA
    with eng.begin() as c:
        c.execute(text(f"DROP SCHEMA IF EXISTS {SRC_SCHEMA} CASCADE"))


def test_postgres_connector_end_to_end(client, workspace, source_table):
    from app.db import get_engine

    ws = workspace["id"]
    cfg = {"host": "localhost", "port": 5432, "database": "lumen_analytics_test", "user": "dw_owner", "schema": source_table}
    r = client.post(api(ws, "/connectors"), json={"name": "Local PG", "kind": "postgresql", "config": cfg})
    assert r.status_code == 201, r.text
    conn = r.json()
    assert conn["config"]["has_saved_password"] is False
    cid = conn["id"]
    assert client.post(api(ws, f"/connectors/{cid}/test")).json()["status"] == "ok"
    assert client.get(api(ws, "/connectors")).json()[0]["last_status"] == "ok"

    tables = client.get(api(ws, f"/connectors/{cid}/tables")).json()
    assert {"schema": source_table, "table": "items", "type": "table"} in tables
    assert {"schema": source_table, "table": "expensive", "type": "view"} in tables

    prev = client.get(api(ws, f"/connectors/{cid}/preview"), params={"table": "items", "schema": source_table}).json()
    assert prev["columns"] == ["id", "name", "amount", "created"] and len(prev["rows"]) == 3
    r = client.get(api(ws, f"/connectors/{cid}/preview"), params={"table": "pg_authid", "schema": "pg_catalog"})
    assert r.status_code == 400
    r = client.get(api(ws, f"/connectors/{cid}/preview"), params={"table": "items; drop table x", "schema": source_table})
    assert r.status_code == 400

    r = client.post(api(ws, f"/connectors/{cid}/import"), json={"tables": [{"schema": source_table, "table": "items", "name": "Items"}]})
    assert r.status_code == 201, r.text
    ds = r.json()[0]
    assert ds["kind"] == "connector" and ds["row_count"] == 3
    full = _dataset(client, ws, ds["id"])
    assert {c["name"]: c["data_type"] for c in full["columns"]} == {"id": "integer", "name": "text", "amount": "decimal", "created": "date"}
    rows = {r[0]: r for r in _rows(client, ws, ds["id"])["rows"]}
    assert rows[1][2] == 1299.5 and rows[3][3] == "2026-03-01"

    # the source changes -> sync picks it up (and recorded prep steps are replayed)
    _step(client, ws, ds["id"], "change_case", {"column": "name", "mode": "upper"})
    with get_engine().begin() as c:
        c.execute(text(f"INSERT INTO {source_table}.items VALUES (4, 'Tray', 99.99, '2026-04-01')"))
        c.execute(text(f"UPDATE {source_table}.items SET amount = 1000 WHERE id = 1"))
    r = client.post(api(ws, f"/datasets/{ds['id']}/sync"))
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 4 and r.json()["dataset"]["row_count"] == 4
    rows = {r[0]: r for r in _rows(client, ws, ds["id"])["rows"]}
    assert rows[1][2] == 1000 and rows[4][1] == "TRAY"

    # the source gains a column (e.g. a warehouse view extended) -> the next sync adds it, typed from the source
    with get_engine().begin() as c:
        c.execute(text(f"ALTER TABLE {source_table}.items ADD COLUMN document_type text, ADD COLUMN shipped date"))
        c.execute(text(f"UPDATE {source_table}.items SET document_type = 'PURCHASE', shipped = '2026-05-01' WHERE id = 1"))
    r = client.post(api(ws, f"/datasets/{ds['id']}/sync"))
    assert r.status_code == 200, r.text
    cols = {c["name"]: c["data_type"] for c in _dataset(client, ws, ds["id"])["columns"]}
    assert cols["document_type"] == "text" and cols["shipped"] == "date"
    rows = {r[0]: r for r in _rows(client, ws, ds["id"])["rows"]}
    assert rows[1][4:6] == ["PURCHASE", "2026-05-01"] and rows[1][1] == "KADAI"       # prep steps still replayed

    r = client.put(api(ws, f"/datasets/{ds['id']}/schedule"), json={"interval_minutes": 60})
    assert r.status_code == 200 and r.json()["interval_minutes"] == 60
    assert client.get(api(ws, "/schedules")).json()[0]["dataset_id"] == ds["id"]

    file_ds = _import_sales(client, ws)["dataset"]
    assert client.post(api(ws, f"/datasets/{file_ds['id']}/sync")).status_code == 400
    assert client.put(api(ws, f"/datasets/{file_ds['id']}/schedule"), json={"interval_minutes": 60}).status_code == 400

    # deleting the connector: sync then explains what happened
    assert client.delete(api(ws, f"/connectors/{cid}")).status_code == 204
    r = client.post(api(ws, f"/datasets/{ds['id']}/sync"))
    assert r.status_code == 400 and "deleted" in r.json()["detail"]


def test_connector_config_sanitizing(client, workspace, monkeypatch):
    ws = workspace["id"]
    monkeypatch.delenv("LUMEN_SECRET_KEY", raising=False)
    r = client.post(api(ws, "/connectors"), json={"name": "x", "kind": "sqlalchemy", "config": {"url": "postgresql://bob:s3cret@db.example.com/sales"}})
    assert r.status_code == 400 and "password" in r.json()["detail"].lower()
    r = client.post(api(ws, "/connectors"), json={"name": "x", "kind": "postgresql", "config": {"host": "h", "password": "s3cret"}})
    assert r.status_code == 400 and "LUMEN_SECRET_KEY" in r.json()["detail"]
    r = client.post(api(ws, "/connectors"), json={"name": "x", "kind": "mongodb", "config": {}})
    assert r.status_code == 400
    r = client.post(api(ws, "/connectors"), json={"name": "x", "kind": "sqlalchemy",
                                                  "config": {"url": "mysql+pymysql://bob@db.example.com/sales", "password_env": "LUMEN_CONN_MY_PW", "evil": "1"}})
    assert r.status_code == 201
    assert r.json()["config"] == {"url": "mysql+pymysql://bob@db.example.com/sales", "password_env": "LUMEN_CONN_MY_PW", "has_saved_password": False}

    from cryptography.fernet import Fernet

    monkeypatch.setenv("LUMEN_SECRET_KEY", Fernet.generate_key().decode())
    r = client.post(api(ws, "/connectors"), json={"name": "enc", "kind": "postgresql", "config": {"host": "h", "user": "u", "password": "s3cret"}})
    assert r.status_code == 201
    assert "s3cret" not in r.text and r.json()["config"]["has_saved_password"] is True
    listing = client.get(api(ws, "/connectors")).text
    assert "s3cret" not in listing and "password_encrypted" not in listing
    from app.db import SessionLocal
    from app import models as m

    with SessionLocal() as db:
        stored = db.get(m.Connector, r.json()["id"]).config
    assert "password" not in stored and stored["password_encrypted"] and "s3cret" not in json.dumps(stored)
