import json, os, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
S = os.path.dirname(os.path.abspath(__file__)); M = json.load(open(f"{S}/_mapping_model.json"))
cov = pd.read_csv(f"{S}/_coverage_report.csv")
F="Arial"; HDR=PatternFill("solid",fgColor="1F3864"); INP=PatternFill("solid",fgColor="FFF2CC"); OK=PatternFill("solid",fgColor="E2EFDA")
thin=Side(style="thin",color="BFBFBF")
wb=Workbook()
def sheet(title, headers, rows, widths, input_cols=(), conf_col=None):
    ws = wb.create_sheet(title); ws.append(headers)
    for i,w in enumerate(widths,1):
        c=ws.cell(row=1,column=i); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=HDR
        c.alignment=Alignment(wrap_text=True,vertical="center"); ws.column_dimensions[c.column_letter].width=w
    for r in rows: ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name=F,size=9); c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(bottom=thin)
            if c.column in input_cols: c.fill=INP
            if conf_col and c.column==conf_col and c.value=="Yes": c.fill=OK
    ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=32
    return ws
yn = lambda x: "Yes" if x else "No"
# Start here
st = wb.active; st.title="Start here"
lines = [("Reference mappings - draft v1.2 for review (adds the virtual locations of the v2 design)",True),("",False),
 ("What this is",True),("Every location and sales-channel name found in the sample files, mapped to one standard code. These tables decide where stock sits and which channel every sale belongs to.",False),
 ("",False),("How it was built",True),
 ("Mappings come from evidence in the data, not name-guessing: Retail Sales links each long contact name to its ERP location; every inward stock transfer links to its outward document, which reveals what each document prefix means (100% consistent in all 21 prefixes).",False),
 ("Coverage: 259 distinct values found across all files. 253 are mapped, 6 need your answer (sheet 'Open questions'), 0 are unaccounted for.",False),
 ("",False),("How to review",True),
 ("1. Answer the 6 items in 'Open questions'.",False),
 ("2. On 'Locations', check Type, Active and Sellable for each row, especially the ones with a question. Write corrections in the yellow column.",False),
 ("3. Skim 'Channels' and 'Channel spellings'. Write corrections in the yellow column.",False),
 ("4. 'Confirmed' = Yes means you already confirmed it earlier. Everything else is a proposal until you approve it.",False),
 ("",False),("What changes after review",True),
 ("Your corrections are applied to the seed CSVs and re-loaded. Approved rows become confirmed. The coverage check is re-run on every future data pull, so any new location or channel name is caught instead of silently going to 'Unknown'.",False)]
for t,b in lines: st.append([t])
for i,(t,b) in enumerate(lines,1):
    c=st.cell(row=i,column=1); c.font=Font(name=F,size=12 if i==1 else 10,bold=b); c.alignment=Alignment(wrap_text=True,vertical="top")
st.column_dimensions["A"].width=120
# Open questions
oq = [(u[2], u[0], u[3], u[4], "") for u in M["UNRESOLVED"]]
extra = [(r[1] or r[2], "Location attributes", f"Stock on 15-Sep: {r[12] if r[12] is not None else 'n/a'} units", r[14], "") for r in M["L"] if r[14]]
sheet("Open questions",["Value / location","Where seen","Evidence","Question","Your answer"], oq+extra,[40,26,40,60,40],input_cols=(5,))
# Locations
rows=[(r[0],r[1] or "(not an ERP location)",r[2],r[3],r[4],r[5],r[6],yn(r[10]),yn(r[8]),r[9],r[12] if r[12] is not None else "",yn(r[11]),r[13],r[14],"") for r in M["L"]]
sheet("Locations",["Code","ERP name","Name","Type","Operated by","GST state","City","Active","Sellable stock","Stock system of record","Trading units 15-Sep","Confirmed","Evidence","Question","Your correction"],
      rows,[22,26,28,20,11,8,12,8,9,12,10,9,50,40,30],input_cols=(15,),conf_col=12)
# Location spellings
sheet("Location spellings",["Source file","Field","Value as it appears","Maps to location","Confirmed","Evidence","Your correction"],
      [(a[0],a[1],a[2],a[3],yn(a[4]),a[5],"") for a in M["LA"]],[24,20,46,22,9,55,28],input_cols=(7,),conf_col=5)
# Transfer prefixes
sheet("Transfer prefixes",["Outward document prefix","Sending location","Linked receipts checked","Evidence","Your correction"],
      [(p,c,n,"Every linked inward receipt names this location as the sender","") for p,c,n in M["PFX"]],[22,24,14,55,28],input_cols=(5,))
# Channels
sheet("Channels",["Code","Channel","Sub-channel","Channel group","Confirmed","Notes","Your correction"],
      [(c[0],c[1],c[2],c[3],yn(c[4]),c[5],"") for c in M["C"]],[26,18,26,16,9,60,28],input_cols=(7,),conf_col=5)
# Channel spellings
cs=[(a[0],a[1],a[2],a[3],yn(a[5]),a[6],"") for a in M["CA"]]
cs+= [("ERP_DOMESTIC_SALES / RETURN","partner type",p,c,yn(k),"Domestic channel comes from the partner type","") for p,c,k in M["PTC"]]
sheet("Channel spellings",["Source file","Field","Value as it appears","Maps to channel","Confirmed","Evidence","Your correction"],cs,[24,16,42,24,9,50,28],input_cols=(7,),conf_col=5)
# Coverage
sheet("Coverage check",["Kind","Source file","Field","Value","Status"],cov.values.tolist(),[14,24,22,46,16])
out="/mnt/user-data/outputs/Reference_Mappings_Draft_v1.2.xlsx"; wb.save(out); print("saved", out)
