"""
Loader test suite.

Pure tests (no database) always run. Database tests need a DISPOSABLE database
built with build_db.sh and pointed to by DW_TEST_DATABASE (plus normal PG* env):
    createdb dw_test && ./build_db.sh dw_test
    DW_TEST_DATABASE=dw_test pytest -q tests/test_loader.py
Every database test removes the batches it created.
Sample-file tests need DW_SAMPLE_DIR (folder with the real sample exports).
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from loader.readers import ReaderError, cell_to_text, read_workbook
from loader.validate import ContractError, load_contract, parse, parse_business_date, parse_stock_pivot

DB = os.environ.get("DW_TEST_DATABASE")
SAMPLES = Path(os.environ["DW_SAMPLE_DIR"]) if os.environ.get("DW_SAMPLE_DIR") else None
needs_db = pytest.mark.skipif(not DB, reason="set DW_TEST_DATABASE to a disposable database")
needs_samples = pytest.mark.skipif(not SAMPLES, reason="set DW_SAMPLE_DIR to the sample export folder")


# ------------------------------------------------------------------ cell conversion
@pytest.mark.parametrize("value,expected", [
    (None, None), ("", None), ("  keep spaces ", "  keep spaces "), (True, "TRUE"), (False, "FALSE"),
    (42, "42"), (124118580.0, "124118580"), (705.93, "705.93"), (0.1 + 0.2, "0.30000000000000004"),
    (-3.0, "-3"), (dt.datetime(2026, 4, 1, 7, 28), "2026-04-01 07:28:00"), (dt.date(2026, 9, 15), "2026-09-15"),
    ("08/04/2026;08/04/2026", "08/04/2026;08/04/2026"),
])
def test_cell_to_text_is_exact(value, expected):
    assert cell_to_text(value) == expected


def test_nan_is_refused():
    with pytest.raises(ReaderError):
        cell_to_text(float("nan"))


@pytest.mark.parametrize("text,fmt,expected", [
    ("2026-04-01 00:00:00", "excel_datetime", dt.date(2026, 4, 1)),
    ("2026-09-02", "text_ymd", dt.date(2026, 9, 2)),
    ("08/04/2026;08/04/2026", "text_dmy_repeated", dt.date(2026, 4, 8)),
    ("22/11/2024;22/01/2025", "text_dmy_repeated", None),              # different dates: not one day
    ("31/02/2026", "text_dmy_repeated", None),                          # impossible date
    ("4/14/2025 15:56", "text_mdy_time", dt.date(2025, 4, 14)),        # month first (WMS)
    ("2026-04-06 00:00:00", "excel_datetime_or_text_dmy_repeated", dt.date(2026, 4, 6)),
    ("06/04/2026", "excel_datetime_or_text_dmy_repeated", dt.date(2026, 4, 6)),
])
def test_business_dates(text, fmt, expected):
    assert parse_business_date(text, fmt) == expected


# ------------------------------------------------------------------ synthetic files
def make_xlsx(path: Path, headers, rows, sheet="Sheet1", extra_sheets=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for name, content in (extra_sheets or {}).items():
        s = wb.create_sheet(name)
        for r in content:
            s.append(r)
    wb.save(path)
    return path


def transfer_file(tmp_path, name="t.xlsx", mutate=None, rows=3, **kw):
    c = load_contract("ERP_STOCK_TRANSFER_OUT")
    headers = [x["header"] for x in c["columns"]]
    data = []
    for i in range(rows):
        r = [None] * len(headers)
        r[headers.index("DOCNUM")] = f"TEST/DC/{i}/26-27"
        r[headers.index("PRODNUM")] = "GSTEA2724"
        r[headers.index("DOCOPENDATE")] = dt.datetime(2026, 4, 1 + i)
        r[headers.index("PRODTRANSQTY")] = -2
        data.append(r)
    if mutate:
        headers, data = mutate(headers, data)
    return make_xlsx(tmp_path / name, headers, data, **kw)


def test_valid_synthetic_file_parses(tmp_path):
    sheet, p = parse(read_workbook(transfer_file(tmp_path)), load_contract("ERP_STOCK_TRANSFER_OUT"))
    assert p.rows_read == 3 and (p.period_start, p.period_end) == (dt.date(2026, 4, 1), dt.date(2026, 4, 3))


@pytest.mark.parametrize("label,mutate,message", [
    ("missing column", lambda h, d: (h[:-1], [r[:-1] for r in d]), "missing"),
    ("new column", lambda h, d: (h + ["NEW COL"], [r + ["x"] for r in d]), "new / unexpected"),
    ("renamed column", lambda h, d: (["docnum" if x == "DOCNUM" else x for x in h], d), "missing"),
    ("duplicate column", lambda h, d: (h + ["DOCNUM"], [r + ["x"] for r in d]), "duplicated"),
    ("value beyond last header", lambda h, d: (h, [r + ["stray"] if i == 1 else r for i, r in enumerate(d)]), "beyond the last header"),
    ("no data rows", lambda h, d: (h, []), "no data rows"),
    ("no parseable dates", lambda h, d: (h, [[("n/a" if h[j] == "DOCOPENDATE" else v) for j, v in enumerate(r)] for r in d]), "no parseable business dates"),
])
def test_contract_violations_are_rejected(tmp_path, label, mutate, message):
    f = transfer_file(tmp_path, mutate=mutate)
    with pytest.raises(ContractError, match=message):
        parse(read_workbook(f), load_contract("ERP_STOCK_TRANSFER_OUT"))


def test_column_order_change_is_accepted(tmp_path):
    f = transfer_file(tmp_path, mutate=lambda h, d: (list(reversed(h)), [list(reversed(r)) for r in d]))
    _, p = parse(read_workbook(f), load_contract("ERP_STOCK_TRANSFER_OUT"))
    i = p.columns.index("docnum") + 1
    assert [r[i] for r in p.records] == ["TEST/DC/0/26-27", "TEST/DC/1/26-27", "TEST/DC/2/26-27"]


def test_wrong_sheet_name_rejected(tmp_path):
    with pytest.raises(ContractError, match="not found"):
        parse(read_workbook(transfer_file(tmp_path, sheet="Data")), load_contract("ERP_STOCK_TRANSFER_OUT"))


def test_data_on_another_sheet_rejected(tmp_path):
    f = transfer_file(tmp_path, extra_sheets={"Sheet2": [["hidden", "rows"]]})
    with pytest.raises(ContractError, match="would be silently lost"):
        parse(read_workbook(f), load_contract("ERP_STOCK_TRANSFER_OUT"))


def test_empty_extra_sheets_allowed(tmp_path):
    f = transfer_file(tmp_path, extra_sheets={"Sheet2": [], "Sheet3": []})
    parse(read_workbook(f), load_contract("ERP_STOCK_TRANSFER_OUT"))


def _instamart(tmp_path, counter):
    c = load_contract("INSTAMART_SALES")
    headers = [None] + [x["header"] for x in c["columns"]]
    rows = [[counter[i]] + [dt.datetime(2026, 9, 1) if h == "ORDERED_DATE" else "x" for h in headers[1:]] for i in range(3)]
    return make_xlsx(tmp_path / "im.xlsx", headers, rows, sheet="Sales Report")


def test_unnamed_row_counter_is_dropped_when_it_is_a_counter(tmp_path):
    _, p = parse(read_workbook(_instamart(tmp_path, [0, 1, 2])), load_contract("INSTAMART_SALES"))
    assert p.rows_read == 3 and "brand" == p.columns[0]


def test_unnamed_column_with_real_data_is_rejected(tmp_path):
    with pytest.raises(ContractError, match="row counter"):
        parse(read_workbook(_instamart(tmp_path, [0, 7, 2])), load_contract("INSTAMART_SALES"))


# ------------------------------------------------------------------ stock report (pure, on the real file)
@pytest.fixture(scope="module")
def stock_rows():
    if not SAMPLES:
        pytest.skip("set DW_SAMPLE_DIR")
    return read_workbook(SAMPLES / "Inventory_report_15_Sept.xls")["ChannelwiseStock"]


def _find(rows, text):
    return next((i, j) for i, r in enumerate(rows) for j, v in enumerate(r) if v == text)


@needs_samples
def test_stock_report_parses(stock_rows):
    p = parse_stock_pivot(stock_rows, load_contract("ERP_INVENTORY_SNAPSHOT"))
    assert p.as_of_in_file == dt.date(2026, 9, 15) and p.rows_read == 4162 and len(p.records) == 4162 * 37
    assert any(w[0] == "stock_report_locations_without_column" for w in p.warnings)


def _bump(rows, r, c, by):
    rows[r][c] = str(int(float((rows[r][c] or "0").replace(",", ""))) + by)


@needs_samples
def test_stock_report_row_not_adding_up_rejected(stock_rows):
    # quantity +5 AND that location's grand total +5: only the per-row Total check can catch it
    rows = copy.deepcopy(stock_rows)
    hr, col = _find(rows, "08 Jaipur Branch")
    footer = max(i for i, r in enumerate(rows) if any(v is not None for v in r))
    _bump(rows, hr + 1, col, 5); _bump(rows, footer, col, 5)
    with pytest.raises(ContractError, match="sum of location quantities != Total"):
        parse_stock_pivot(rows, load_contract("ERP_INVENTORY_SNAPSHOT"))


@needs_samples
def test_stock_report_column_not_adding_up_rejected(stock_rows):
    # quantity +5 AND the row Total +5: only the grand-total (column) check can catch it
    rows = copy.deepcopy(stock_rows)
    hr, col = _find(rows, "08 Jaipur Branch")
    _, tcol = _find(rows[hr:hr + 1], "Total")
    _bump(rows, hr + 1, col, 5); _bump(rows, hr + 1, tcol, 5)
    with pytest.raises(ContractError, match="grand-total row"):
        parse_stock_pivot(rows, load_contract("ERP_INVENTORY_SNAPSHOT"))


@needs_samples
def test_stock_report_truncated_rejected(stock_rows):
    rows = copy.deepcopy(stock_rows)
    footer = max(i for i, r in enumerate(rows) if any(v is not None for v in r))
    with pytest.raises(ContractError, match="no grand-total row"):
        parse_stock_pivot(rows[:footer], load_contract("ERP_INVENTORY_SNAPSHOT"))


@needs_samples
def test_stock_report_dropped_product_row_rejected(stock_rows):
    rows = copy.deepcopy(stock_rows)
    hr, _ = _find(rows, "Prod Num")
    del rows[hr + 1]                                                 # footer totals no longer match
    with pytest.raises(ContractError, match="grand-total"):
        parse_stock_pivot(rows, load_contract("ERP_INVENTORY_SNAPSHOT"))


@needs_samples
def test_stock_report_multi_day_rejected(stock_rows):
    rows = copy.deepcopy(stock_rows)
    rows[2][0] = "9/14/2026 To 9/15/2026"
    with pytest.raises(ContractError, match="single day"):
        parse_stock_pivot(rows, load_contract("ERP_INVENTORY_SNAPSHOT"))


# ------------------------------------------------------------------ database behaviour
@pytest.fixture
def db():
    import psycopg
    conn = psycopg.connect(os.environ.get("DATABASE_URL", ""), dbname=DB, autocommit=False)
    start = conn.execute("SELECT coalesce(max(batch_id),0) FROM meta.load_batch").fetchone()[0]
    conn.commit()
    yield conn
    conn.rollback()
    tables = [r[0] for r in conn.execute("SELECT table_name FROM information_schema.tables WHERE table_schema='raw'")]
    for t in tables:
        conn.execute(f"DELETE FROM raw.{t} WHERE batch_id > %s", (start,))
    conn.execute("DELETE FROM meta.dq_check_result WHERE batch_id > %s", (start,))
    conn.execute("DELETE FROM meta.load_batch WHERE batch_id > %s", (start,))
    conn.commit()
    conn.close()


def _batches(db, since_file):
    rows = db.execute("SELECT status, rows_loaded, notes FROM meta.load_batch WHERE file_name=%s ORDER BY batch_id", (since_file,)).fetchall()
    db.rollback()
    return rows


def _raw_count(db, docnum_prefix="TEST/"):
    n = db.execute("SELECT count(*) FROM raw.erp_stock_transfer_out WHERE docnum LIKE %s", (docnum_prefix + "%",)).fetchone()[0]
    db.rollback()
    return n


@needs_db
def test_load_commits_rows_and_audit(tmp_path, db):
    from loader.load import load_file
    o = load_file("ERP_STOCK_TRANSFER_OUT", transfer_file(tmp_path, "ok.xlsx"), loaded_by="pytest", conn=db)
    assert o.rows_loaded == 3 and _raw_count(db) == 3
    assert _batches(db, "ok.xlsx") == [("LOADED", 3, None)]


@needs_db
def test_same_file_twice_is_skipped(tmp_path, db):
    from loader.load import AlreadyLoaded, load_file
    f = transfer_file(tmp_path, "twice.xlsx")
    load_file("ERP_STOCK_TRANSFER_OUT", f, loaded_by="pytest", conn=db)
    with pytest.raises(AlreadyLoaded):
        load_file("ERP_STOCK_TRANSFER_OUT", f, loaded_by="pytest", conn=db)
    assert _raw_count(db) == 3


@needs_db
def test_rejected_file_is_audited_and_lands_nothing(tmp_path, db):
    from loader.load import load_file
    f = transfer_file(tmp_path, "bad.xlsx", mutate=lambda h, d: (h[:-1], [r[:-1] for r in d]))
    with pytest.raises(ContractError):
        load_file("ERP_STOCK_TRANSFER_OUT", f, loaded_by="pytest", conn=db)
    (status, rows, notes), = _batches(db, "bad.xlsx")
    assert status == "FAILED" and rows == 0 and "missing" in notes and _raw_count(db) == 0


@needs_db
def test_failure_mid_load_rolls_back_everything(tmp_path, db, monkeypatch):
    import loader.load as L
    real = L._load

    def boom(conn, contract, parsed, *a):
        real(conn, contract, parsed, *a)           # rows are copied in...
        raise RuntimeError("simulated crash")      # ...then the process fails before commit
    monkeypatch.setattr(L, "_load", boom)
    with pytest.raises(RuntimeError):
        L.load_file("ERP_STOCK_TRANSFER_OUT", transfer_file(tmp_path, "crash.xlsx"), loaded_by="pytest", conn=db)
    assert _raw_count(db) == 0
    assert [b[0] for b in _batches(db, "crash.xlsx")] == ["FAILED"]


@needs_db
def test_failed_file_can_be_retried(tmp_path, db, monkeypatch):
    import loader.load as L
    f = transfer_file(tmp_path, "retry.xlsx")
    monkeypatch.setattr(L, "_load", lambda *a: (_ for _ in ()).throw(RuntimeError("disk full")))
    with pytest.raises(RuntimeError):
        L.load_file("ERP_STOCK_TRANSFER_OUT", f, loaded_by="pytest", conn=db)
    monkeypatch.undo()
    L.load_file("ERP_STOCK_TRANSFER_OUT", f, loaded_by="pytest", conn=db)
    assert [b[0] for b in _batches(db, "retry.xlsx")] == ["FAILED", "LOADED"] and _raw_count(db) == 3


@needs_db
def test_policy_refusals(tmp_path, db):
    from loader.load import LoadRefused, load_file
    f = transfer_file(tmp_path)
    with pytest.raises(LoadRefused, match="Do not pass --as-of"):
        load_file("ERP_STOCK_TRANSFER_OUT", f, as_of=dt.date(2026, 9, 15), conn=db)
    with pytest.raises(LoadRefused, match="--as-of YYYY-MM-DD is required"):
        load_file("BLINKIT_INVENTORY", f, conn=db)
    with pytest.raises(LoadRefused, match="DRAFT"):
        load_file("WMS_DISPATCH", f, conn=db)
    with pytest.raises(ContractError, match="no contract"):
        load_file("NOT_A_SOURCE", f, conn=db)


@needs_db
@needs_samples
def test_state_export_cannot_contain_future_dates(db):
    from loader.load import load_file
    with pytest.raises(ContractError, match="after --as-of"):
        load_file("ERP_PURCHASE_ORDER", SAMPLES / "Pending_Purchase_Order_26-27_12_09.xlsx", as_of=dt.date(2026, 9, 12), conn=db)
